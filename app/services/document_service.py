import os
import shutil
import uuid
from typing import List, Optional
from app.core.config import settings
from app.core.database import get_db
from sqlalchemy.orm import Session
from app.models.document import Document
from app.services.rag_service import RAGService
import PyPDF2
import docx
from datetime import datetime

UPLOAD_DIR = "uploads"

class DocumentService:
    def __init__(self, db: Session = None):
        self.db = db or next(get_db())
        self.rag_service = RAGService(self.db)

    def _save_upload_file(self, file) -> str:
        """Save the uploaded file to the uploads directory and return the relative path."""
        # Ensure the upload directory exists
        os.makedirs(UPLOAD_DIR, exist_ok=True)
        # Generate a unique filename to avoid collisions
        import uuid
        file_extension = os.path.splitext(file.filename)[1]
        unique_filename = f"{uuid.uuid4()}{file_extension}"
        file_path = os.path.join(UPLOAD_DIR, unique_filename)
        # Save the file
        with open(file_path, "wb") as buffer:
            # If the file has a 'file' attribute (like UploadFile), use it, ...)
            buffer.write(file.file.read() if hasattr(file, 'file') else file.read())
        return file_path

    def _extract_text_from_pdf(self, file_path: str) -> str:
        """Extract text from a PDF file."""
        text = ""
        with open(file_path, "rb") as file:
            pdf_reader = PyPDF2.PdfReader(file)
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
        return text

    def _extract_text_from_docx(self, file_path: str) -> str:
        """Extract text from a DOCX file."""
        doc = docx.Document(file_path)
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        return text

    async def upload_document(self, user_id: str, tenant_id: str, file: UploadFile) -> dict:
        # Save the uploaded file
        storage_path = self._save_upload_file(file)
        # Extract text based on file type
        extracted_text = ""
        if file.content_type == "text/plain":
            with open(storage_path, "r", encoding="utf-8") as f:
                extracted_text = f.read()
        elif file.content_type == "application/pdf":
            extracted_text = self._extract_text_from_pdf(storage_path)
        elif file.content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            extracted_text = self._extract_text_from_docx(storage_path)
        else:
            # For unsupported file types, we can either raise an error or skip text extraction
            # For now, we'll skip text extraction and set extracted_text to empty
            extracted_text = ""
        # Create document record
        document = Document(
            id=str(uuid.uuid4()),
            user_id=user_id,
            tenant_id=tenant_id,
            filename=file.filename,
            content_type=file.content_type,
            size=os.path.getsize(storage_path),
            storage_path=storage_path,
            extracted_text=extracted_text,
            processing_status="completed" if extracted_text else "failed",
        )
        self.db.add(document)
        self.db.commit()
        self.db.refresh(document)
        # Process document with RAG
        await self.rag_service.process_document(document.id, document.extracted_text, tenant_id)
        return {
            "id": document.id,
            "filename": document.filename,
            "content_type": document.content_type,
            "size": document.size,
            "created_at": document.created_at.isoformat() if document.created_at else None,
            "updated_at": document.updated_at.isoformat() if document.updated_at else None,
        }

    async def list_documents(self, user_id: str, tenant_id: str) -> List[dict]:
        documents = self.db.query(Document).filter(
            Document.user_id == user_id,
            Document.tenant_id == tenant_id
        ).all()
        return [
            {
                "id": doc.id,
                "filename": doc.filename,
                "content_type": doc.content_type,
                "size": doc.size,
                "created_at": doc.created_at.isoformat() if doc.created_at else None,
                "updated_at": doc.updated_at.isoformat() if doc.updated_at else None,
            }
            for doc in documents
        ]

    async def get_document(self, document_id: str, user_id: str, tenant_id: str) -> Optional[dict]:
        document = self.db.query(Document).filter(
            Document.id == document_id,
            Document.user_id == user_id,
            Document.tenant_id == tenant_id
        ).first()
        if not document:
            return None
        return {
            "id": document.id,
            "filename": document.filename,
            "content_type": document.content_type,
            "size": document.size,
            "created_at": document.created_at.isoformat() if document.created_at else None,
            "updated_at": document.updated_at.isoformat() if document.updated_at else None,
        }

    async def delete_document(self, document_id: str, user_id: str, tenant_id: str) -> bool:
        document = self.db.query(Document).filter(
            Document.id == document_id,
            Document.user_id == user_id,
            Document.tenant_id == tenant_id
        ).first()
        if not document:
            return False
        # Delete file from disk
        if os.path.exists(document.storage_path):
            os.remove(document.storage_path)
        self.db.delete(document)
        self.db.commit()
        return True

